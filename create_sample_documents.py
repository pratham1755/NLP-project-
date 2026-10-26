"""
Generates real sample PDF and DOCX files for testing file uploads in Streamlit.
"""

import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document

base_dir = os.path.dirname(os.path.abspath(__file__))
sample_dir = os.path.join(base_dir, "sample_data")
txt_resume_path = os.path.join(sample_dir, "sample_data_analyst_resume.txt")

with open(txt_resume_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

# 1. Generate PDF
pdf_path = os.path.join(sample_dir, "sample_data_analyst_resume.pdf")
doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
styles = getSampleStyleSheet()

story = []
for line in lines:
    line_clean = line.strip()
    if not line_clean:
        story.append(Spacer(1, 8))
    elif line_clean.startswith("John Doe"):
        title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=16, leading=20)
        story.append(Paragraph(line_clean, title_style))
    elif line_clean.isupper() and len(line_clean) < 30:
        h2_style = ParagraphStyle('H2Style', parent=styles['Heading2'], fontSize=12, leading=16, spaceBefore=6)
        story.append(Paragraph(line_clean, h2_style))
    else:
        body_style = ParagraphStyle('BodyStyle', parent=styles['Normal'], fontSize=9, leading=12)
        # Escape XML entities for ReportLab
        clean_xml = line_clean.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(clean_xml, body_style))

doc.build(story)
print(f"[OK] Created {pdf_path}")

# 2. Generate DOCX
docx_path = os.path.join(sample_dir, "sample_data_analyst_resume.docx")
docx_doc = Document()

for line in lines:
    line_clean = line.strip()
    if not line_clean:
        continue
    elif line_clean.startswith("John Doe"):
        docx_doc.add_heading(line_clean, level=1)
    elif line_clean.isupper() and len(line_clean) < 30:
        docx_doc.add_heading(line_clean, level=2)
    elif line_clean.startswith("•"):
        docx_doc.add_paragraph(line_clean[1:].strip(), style='List Bullet')
    else:
        docx_doc.add_paragraph(line_clean)

docx_doc.save(docx_path)
print(f"[OK] Created {docx_path}")
