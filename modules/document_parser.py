"""
Module 1: Resume Upload and Parsing
Extracts raw text from PDF, DOCX, and TXT resumes with multi-library fallback support.
"""

import io
import re
from typing import Tuple, Optional


def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, Optional[str]]:
    """
    Extracts text from a PDF file byte stream using PyMuPDF (fitz),
    pdfminer, or PyPDF2 as fallback.
    Returns (extracted_text, warning_or_error_message).
    """
    text = ""
    error_msg = None

    # Method 1: Try PyMuPDF (fitz)
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages_text = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            pages_text.append(page.get_text())
        text = "\n".join(pages_text).strip()
        if text:
            return text, None
    except Exception:
        pass

    # Method 2: Try pdfminer.six
    try:
        from pdfminer.high_level import extract_text as pdfminer_extract
        text = pdfminer_extract(io.BytesIO(file_bytes)).strip()
        if text:
            return text, None
    except Exception:
        pass

    # Method 3: Try PyPDF2
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                pages_text.append(extracted)
        text = "\n".join(pages_text).strip()
        if text:
            return text, None
    except Exception as e:
        error_msg = f"Failed to extract PDF text: {str(e)}"

    if not text:
        error_msg = (
            "No extractable text found. The PDF might be scanned or image-based. "
            "Please upload a text-based PDF or copy/paste the text directly."
        )

    return text, error_msg


def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, Optional[str]]:
    """
    Extracts text from a DOCX file byte stream using python-docx with docx2txt fallback.
    """
    text = ""
    error_msg = None

    # Method 1: Try python-docx
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
        text = "\n".join(full_text).strip()
        if text:
            return text, None
    except Exception:
        pass

    # Method 2: Try docx2txt fallback
    try:
        import docx2txt
        text = docx2txt.process(io.BytesIO(file_bytes)).strip()
        if text:
            return text, None
    except Exception as e:
        error_msg = f"Failed to extract DOCX text: {str(e)}"

    if not text and not error_msg:
        error_msg = "Could not extract text from the DOCX file. Ensure the file contains readable text."

    return text, error_msg


def extract_text_from_txt(file_bytes: bytes) -> Tuple[str, Optional[str]]:
    """
    Extracts text from plain text bytes trying multiple encodings.
    """
    encodings = ["utf-8", "latin-1", "cp1252", "ascii"]
    for enc in encodings:
        try:
            return file_bytes.decode(enc).strip(), None
        except UnicodeDecodeError:
            continue
    return "", "Failed to decode text file. Ensure it is encoded in standard UTF-8 or ASCII."


def parse_resume_file(uploaded_file) -> Tuple[str, Optional[str]]:
    """
    Universal parser for Streamlit uploaded file object.
    Supports PDF, DOCX, and TXT.
    Returns (cleaned_text, warning_message).
    """
    if uploaded_file is None:
        return "", "No file uploaded."

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        text, err = extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        text, err = extract_text_from_docx(file_bytes)
    elif filename.endswith(".txt"):
        text, err = extract_text_from_txt(file_bytes)
    else:
        return "", f"Unsupported file extension: {uploaded_file.name}. Please upload PDF, DOCX, or TXT."

    if text:
        # Standardize extra newlines and carriage returns
        text = re.sub(r"\r\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = text.strip()

    return text, err
